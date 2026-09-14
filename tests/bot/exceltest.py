# -*- coding: utf-8 -*-
"""Сводка Excel: критичные значения должны быть окрашены."""
import asyncio
from datetime import datetime, timedelta
from openpyxl import load_workbook
from database import db


async def add(uuid, h, mb_in, mb_out, packets, peak):
    await db.execute(
        "INSERT INTO traffic_hourly (user_uuid, hour, bytes_in, bytes_out,"
        " packets_in, packets_out, peak_pps, samples) VALUES ($1,"
        " date_trunc('hour', NOW() - ($2||' hours')::INTERVAL),$3,$4,$5,$6,$7,240)"
        " ON CONFLICT (user_uuid, hour) DO NOTHING",
        uuid, str(h), int(mb_in*1048576), int(mb_out*1048576), packets//2, packets//2, peak)


async def main():
    await db.connect()
    await db.execute("DELETE FROM users WHERE uuid LIKE 'xl-%'")
    await db.execute("INSERT INTO users (name, uuid) VALUES ('Торрент','xl-1')")
    await db.execute("INSERT INTO users (name, uuid) VALUES ('Обычный','xl-2')")
    await db.execute("INSERT INTO users (name, uuid, is_active) VALUES ('Пауза','xl-3',FALSE)")
    for h in range(1, 20):
        await add("xl-1", h, 700, 800, 5_000_000, 7200)   # раздача, мелкие пакеты, пик
        await add("xl-2", h, 30, 400, 400_000, 900)       # обычное потребление
    await db.execute(
        "INSERT INTO pps_events (user_uuid, started_at, peak_pps, avg_packet_size)"
        " VALUES ('xl-1', NOW() - INTERVAL '2 hours', 7200, 310),"
        "        ('xl-1', NOW() - INTERVAL '5 hours', 7100, 300),"
        "        ('xl-1', NOW() - INTERVAL '9 hours', 6900, 290)")

    path = "/tmp/summary.xlsx"
    await db.export_summary_to_excel(path, days=30)
    ws = load_workbook(path).active

    rows = {ws.cell(row=r, column=1).value: r for r in range(2, ws.max_row + 1)}
    print("строки:", list(rows))

    def fill(r, c):
        f = ws.cell(row=r, column=c).fill
        return (f.fgColor.rgb or "")[-6:] if f and f.fgColor else ""

    t = rows["Торрент"]
    o = rows["Обычный"]
    p = rows["Пауза"]

    print("торрент: отдача", fill(t, 7), "пакет", fill(t, 8), "пик", fill(t, 9),
          "превышения", fill(t, 10), "вывод", fill(t, 14))
    assert fill(t, 7) == "C00000", "высокая отдача должна быть красной"
    assert fill(t, 8) == "C00000", "мелкие пакеты должны быть красными"
    assert fill(t, 9) == "C00000", "пик у потолка должен быть красным"
    assert fill(t, 10) == "C00000", "три превышения — красное"
    assert fill(t, 14) == "C00000", "вывод про торрент должен быть красным"
    print("критичное окрашено красным: ок")

    for col in (7, 8, 9, 10):
        assert fill(o, col) in ("", "000000", "00000000"), f"обычный не должен краситься (кол. {col}): {fill(o,col)}"
    # вердикт у обычного может быть жёлтым (например, ночная активность),
    # но красным — не должен: красный означает торрент или превышения
    assert fill(o, 14) != "C00000", "обычного нельзя помечать как проблемного"
    print("обычный человек не раскрашен: ок")

    assert fill(t, 3) == "107C41" and fill(p, 3) == "7F7F7F"
    print("статус: активный зелёный, пауза серая — ок")

    await db.execute("DELETE FROM users WHERE uuid LIKE 'xl-%'")
    print("\nВСЁ ПРОШЛО")

asyncio.run(main())
